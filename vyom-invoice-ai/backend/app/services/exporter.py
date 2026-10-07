import json
import csv
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from typing import Dict, Any

def export_json(record: Dict[str, Any]) -> str:
    """Format record as standardized JSON string."""
    return json.dumps(record, indent=2, default=str, ensure_ascii=False)

def export_csv(record: Dict[str, Any]) -> str:
    """Format invoice record into CSV with metadata header and line item table."""
    output = io.StringIO()
    writer = csv.writer(output)

    inv = record.get("invoice") or {}
    supplier = record.get("supplier") or {}
    buyer = record.get("buyer") or {}
    totals = record.get("totals") or {}
    tax = record.get("tax") or {}

    # Header metadata
    writer.writerow(["INVOICE INTELLIGENCE REPORT - VYOM+"])
    writer.writerow(["Invoice Number", inv.get("invoice_number", ""), "Invoice Date", inv.get("invoice_date", "")])
    writer.writerow(["Supplier", supplier.get("name", ""), "Supplier GSTIN", supplier.get("gstin", "")])
    writer.writerow(["Buyer", buyer.get("name", ""), "Buyer GSTIN", buyer.get("gstin", "")])
    writer.writerow(["Place of Supply", inv.get("place_of_supply", ""), "Supply Type", inv.get("supply_type", "")])
    writer.writerow([])

    # Line item columns
    writer.writerow(["#", "Description", "HSN/SAC", "Qty", "Unit", "Rate", "Discount", "Taxable Value", "GST%", "CGST", "SGST", "IGST", "Cess", "Total"])

    for idx, item in enumerate(record.get("items", []), 1):
        writer.writerow([
            idx,
            item.get("description", ""),
            item.get("hsn_sac", ""),
            item.get("quantity", ""),
            item.get("unit", ""),
            item.get("unit_price", ""),
            item.get("discount", 0.0),
            item.get("taxable_value", ""),
            item.get("gst_rate", ""),
            item.get("cgst", ""),
            item.get("sgst", ""),
            item.get("igst", ""),
            item.get("cess", 0.0),
            item.get("total", "")
        ])

    writer.writerow([])
    writer.writerow(["Subtotal", totals.get("subtotal", "")])
    writer.writerow(["Total Tax", tax.get("total_tax", "")])
    writer.writerow(["Grand Total", totals.get("grand_total", "")])

    return output.getvalue()

def export_excel(record: Dict[str, Any]) -> bytes:
    """Generate multi-tab formatted Excel workbook."""
    wb = openpyxl.Workbook()

    # Tab 1: Invoice Overview
    ws1 = wb.active
    ws1.title = "Invoice Summary"
    _write_invoice_sheet(ws1, record)

    # Tab 2: Line Items
    ws2 = wb.create_sheet("Line Items")
    _write_items_sheet(ws2, record)

    # Tab 3: Statutory Validation
    ws3 = wb.create_sheet("Validation")
    _write_validation_sheet(ws3, record)

    # Tab 4: Confidence Scores
    ws4 = wb.create_sheet("Confidence")
    _write_confidence_sheet(ws4, record)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def _write_invoice_sheet(ws, record: dict):
    header_font = Font(bold=True)
    inv = record.get("invoice") or {}
    supplier = record.get("supplier") or {}
    buyer = record.get("buyer") or {}
    totals = record.get("totals") or {}
    tax = record.get("tax") or {}

    rows = [
        ("INVOICE SUMMARY", ""),
        ("Invoice Number", inv.get("invoice_number")),
        ("Invoice Date", inv.get("invoice_date")),
        ("Place of Supply", inv.get("place_of_supply")),
        ("Supply Type", inv.get("supply_type")),
        ("", ""),
        ("SUPPLIER DETAILS", ""),
        ("Name", supplier.get("name")),
        ("GSTIN", supplier.get("gstin")),
        ("State Code", supplier.get("state_code")),
        ("Address", supplier.get("address")),
        ("", ""),
        ("BUYER DETAILS", ""),
        ("Name", buyer.get("name")),
        ("GSTIN", buyer.get("gstin")),
        ("State Code", buyer.get("state_code")),
        ("Address", buyer.get("address")),
        ("", ""),
        ("TOTALS & TAX", ""),
        ("Subtotal", totals.get("subtotal")),
        ("Total Tax", tax.get("total_tax")),
        ("Grand Total", totals.get("grand_total")),
        ("Amount in Words", totals.get("amount_in_words")),
    ]
    for r, (label, value) in enumerate(rows, 1):
        ws.cell(r, 1, label).font = header_font
        ws.cell(r, 2, value or "")

def _write_items_sheet(ws, record: dict):
    headers = ["#", "Description", "HSN/SAC", "Qty", "Unit", "Rate", "Discount", "Taxable Value", "GST%", "CGST", "SGST", "IGST", "Cess", "Total"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill = PatternFill("solid", fgColor="1E3A5F")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")

    for r, item in enumerate(record.get("items", []), 2):
        ws.cell(r, 1, r - 1)
        ws.cell(r, 2, item.get("description"))
        ws.cell(r, 3, item.get("hsn_sac"))
        ws.cell(r, 4, item.get("quantity"))
        ws.cell(r, 5, item.get("unit"))
        ws.cell(r, 6, item.get("unit_price"))
        ws.cell(r, 7, item.get("discount", 0))
        ws.cell(r, 8, item.get("taxable_value"))
        ws.cell(r, 9, item.get("gst_rate"))
        ws.cell(r, 10, item.get("cgst"))
        ws.cell(r, 11, item.get("sgst"))
        ws.cell(r, 12, item.get("igst"))
        ws.cell(r, 13, item.get("cess"))
        ws.cell(r, 14, item.get("total"))

def _write_validation_sheet(ws, record: dict):
    val = record.get("validation") or {}
    ws.cell(1, 1, "Validation Status").font = Font(bold=True)
    ws.cell(1, 2, val.get("status", "unknown"))

    ws.cell(3, 1, "Validation Errors").font = Font(bold=True)
    errors = val.get("errors") or []
    if errors:
        for i, e in enumerate(errors, 4):
            ws.cell(i, 1, f"ERROR #{i-3}")
            ws.cell(i, 2, e)
    else:
        ws.cell(4, 1, "No critical errors found.")

    offset = max(5, 4 + len(errors) + 1)
    ws.cell(offset, 1, "Validation Warnings").font = Font(bold=True)
    warnings = val.get("warnings") or []
    if warnings:
        for i, w in enumerate(warnings, offset + 1):
            ws.cell(i, 1, f"WARNING #{i-offset}")
            ws.cell(i, 2, w)
    else:
        ws.cell(offset + 1, 1, "No warnings found.")

def _write_confidence_sheet(ws, record: dict):
    conf = record.get("confidence") or {}
    ws.cell(1, 1, "Field").font = Font(bold=True)
    ws.cell(1, 2, "Confidence Score").font = Font(bold=True)
    ws.cell(1, 3, "Threshold Assessment").font = Font(bold=True)

    skip = {"overall", "review_required", "low_confidence_fields"}
    r = 2
    for field, value in conf.items():
        if field in skip or not isinstance(value, (int, float)):
            continue
        ws.cell(r, 1, field)
        ws.cell(r, 2, f"{value:.1%}")
        status = "Approved (High)" if value >= 0.90 else ("Warning (Review)" if value >= 0.70 else "Review Required")
        ws.cell(r, 3, status)
        r += 1
